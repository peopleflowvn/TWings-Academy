"""
The default acceptance document (biên bản nghiệm thu + bảng kê), until each client's own template is
known (TMS questions A1, A2): an Excel workbook (export_xlsx) and the same content as a printable A4
page (corporate/acceptance_print.html, "In / Lưu PDF"), both from document(). Three parts:

1. "Biên bản nghiệm thu": parties A (client) and B (the provider, settings.CORPORATE_PROVIDER), the
   contract it rests on, the services and value with the amount in words, quality, payment, copies,
   signatures. Blank details print as dotted lines to fill in by hand.
2. "Bảng kê theo lớp": each class: dates, sessions, hours, learners listed / attended / passed.
3. "Danh sách học viên": the learners of every class from its published results (the evidence).
"""

from io import BytesIO

from django.conf import settings

from .acceptance import totals

DIGITS = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
BLANK = "……………………………………"


def _three(n: int, full: bool) -> str:
    """Read 0-999; `full` = not the leading group, so "không trăm" / "linh" are spoken."""
    hundreds, tens, units = n // 100, n // 10 % 10, n % 10
    words = []
    if hundreds or full:
        words += [DIGITS[hundreds], "trăm"]
    if tens == 0:
        if units and (hundreds or full):
            words.append("linh")
    elif tens == 1:
        words.append("mười")
    else:
        words += [DIGITS[tens], "mươi"]
    if units:
        if units == 1 and tens >= 2:
            words.append("mốt")
        elif units == 5 and tens >= 1:
            words.append("lăm")
        elif units == 4 and tens >= 2:
            words.append("tư")
        else:
            words.append(DIGITS[units])
    return " ".join(words)


def _read(n: int, started: bool) -> str:
    """Read n (any size); `started` = a higher non-zero group was read before ("không trăm", "linh")."""
    if n >= 10**9:
        high, low = divmod(n, 10**9)
        text = _read(high, started) + " tỷ"
        return text + (" " + _read(low, True) if low else "")
    parts = []
    for value, scale in ((n // 10**6, " triệu"), (n // 1000 % 1000, " nghìn"), (n % 1000, "")):
        if value:
            parts.append(_three(value, started) + scale)
            started = True
    return " ".join(parts)


def amount_in_words(amount: int) -> str:
    """17480000 -> "Mười bảy triệu bốn trăm tám mươi nghìn đồng" (negative: "Âm ...")."""
    if amount == 0:
        return "Không đồng"
    text = _read(abs(amount), False)
    text = text[0].upper() + text[1:]
    return ("Âm " + text[0].lower() + text[1:] if amount < 0 else text) + " đồng"


def _money(n) -> str:
    return f"{int(n):,}".replace(",", ".")


def _day(d) -> str:
    return d.strftime("%d/%m/%Y") if d else "…/…/……"


def _class_facts(batch) -> list[dict]:
    """One row per class of the batch (service lines), with its schedule and learner counts."""
    seen, rows = set(), []
    for line in batch.lines.filter(kind="service").select_related(
        "corporate_class__cohort__course", "snapshot"
    ):
        c = line.corporate_class
        if c is None or c.pk in seen:
            continue
        seen.add(c.pk)
        sessions = list(c.cohort.sessions.order_by("starts_at"))
        learners = line.snapshot.learners if line.snapshot else []
        rows.append(
            {
                "class": c,
                "first": sessions[0].starts_at.date() if sessions else c.cohort.start_date,
                "last": sessions[-1].ends_at.date() if sessions else None,
                "sessions": len(sessions),
                "hours": round(sum((s.ends_at - s.starts_at).total_seconds() for s in sessions) / 3600, 2),
                "listed": len(learners),
                "attended": sum(1 for r in learners if r.get("sessions_attended")),
                "passed": sum(1 for r in learners if r.get("result") == "passed"),
                "learners": learners,
            }
        )
    return rows


def document(batch) -> dict:
    """Everything the minutes show, shared by the workbook and the printable page."""
    contract, client, provider = batch.contract, batch.contract.client, settings.CORPORATE_PROVIDER
    t = totals(batch)
    lines = []
    for i, row in enumerate(batch.lines.order_by("sort_order"), start=1):
        quantity = float(row.quantity) if row.quantity != row.quantity.to_integral() else int(row.quantity)
        lines.append(
            {
                "no": i,
                "description": ("[Điều chỉnh] " if row.kind == "adjustment" else "") + row.description,
                "unit": row.unit,
                "quantity": quantity,
                "unit_price": _money(row.unit_price),
                "amount": _money(row.amount),
            }
        )
    period = (
        f"từ {_day(batch.period_from)} đến {_day(batch.period_to)}"
        if batch.period_from or batch.period_to
        else ""
    )
    return {
        "batch": batch,
        "contract": contract,
        "provider": provider,
        "signed": f" ký ngày {_day(contract.signed_on)}" if contract.signed_on else "",
        "period": period,
        "parties": [
            {
                "title": "BÊN A (Bên sử dụng dịch vụ)",
                "name": client.name,
                "address": client.address or BLANK,
                "tax_code": client.tax_code or BLANK,
                "representative": client.representative or BLANK,
                "rep_title": client.representative_title or BLANK,
            },
            {
                "title": "BÊN B (Bên cung cấp dịch vụ)",
                "name": provider["legal_name"],
                "address": provider["address"] or BLANK,
                "phone": provider.get("phone", ""),
                "tax_code": provider["tax_code"] or BLANK,
                "representative": provider["representative"] or BLANK,
                "rep_title": provider["title"] or BLANK,
            },
        ],
        "client": client,
        "lines": lines,
        "totals": {k: _money(v) for k, v in t.items() if k != "lines"},
        "words": amount_in_words(t["total"]),
        "facts": _class_facts(batch),
        "result_labels": {"passed": "Đạt", "not_passed": "Không đạt", "pending": "Chờ xác minh"},
    }


def export_xlsx(batch) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    doc = document(batch)
    contract, client, provider = batch.contract, batch.contract.client, settings.CORPORATE_PROVIDER
    t = totals(batch)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold, center = Font(bold=True), Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    book = Workbook()
    ws = book.active
    ws.title = "Biên bản nghiệm thu"
    for col, width in zip("ABCDEF", (6, 46, 10, 11, 15, 17), strict=True):
        ws.column_dimensions[col].width = width
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    def line(text="", *, merge=True, font=None, align=None, height=None):
        ws.append([text])
        r = ws.max_row
        if merge:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        cell = ws.cell(row=r, column=1)
        cell.font = font or Font()
        cell.alignment = align or wrap
        if height:
            ws.row_dimensions[r].height = height
        return r

    # Header: provider on the left, national motto on the right.
    ws.append([provider["legal_name"].upper(), None, None, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"])
    ws.merge_cells("A1:C1")
    ws.merge_cells("D1:F1")
    ws.append(["Số: ……/BBNT", None, None, "Độc lập – Tự do – Hạnh phúc"])
    ws.merge_cells("A2:C2")
    ws.merge_cells("D2:F2")
    for cell in ("A1", "D1", "D2"):
        ws[cell].font, ws[cell].alignment = bold, center
    ws["A2"].alignment = center
    line()
    line("BIÊN BẢN NGHIỆM THU DỊCH VỤ ĐÀO TẠO", font=Font(bold=True, size=14), align=center, height=22)
    line(f"({batch.title})", align=center)
    line()
    line(
        f"– Căn cứ Hợp đồng dịch vụ đào tạo số {contract.number}{doc['signed']} "
        f"giữa {client.name} và {provider['legal_name']};"
    )
    period = doc["period"]
    line(
        f"– Căn cứ kết quả tổ chức các lớp đào tạo{(' ' + period) if period else ''} "
        "và các tài liệu kèm theo."
    )
    line("Hôm nay, ngày …… tháng …… năm ……, tại ……………………………, chúng tôi gồm:")
    for party in doc["parties"]:
        line(f"{party['title']}: {party['name']}", font=bold)
        line(f"Địa chỉ: {party['address']}")
        if party.get("phone"):
            line(f"Điện thoại: {party['phone']}")
        line(f"Mã số thuế: {party['tax_code']}")
        line(f"Đại diện: {party['representative']}        Chức vụ: {party['rep_title']}")
    line("Hai bên thống nhất nghiệm thu dịch vụ đào tạo với các nội dung sau:")
    line("Điều 1. Nội dung, khối lượng và giá trị nghiệm thu", font=bold)
    header = ["STT", "Nội dung", "ĐVT", "Khối lượng", "Đơn giá (đ)", "Thành tiền (đ)"]
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font, cell.alignment, cell.border = bold, center, box
    for row in doc["lines"]:
        ws.append(
            [row["no"], row["description"], row["unit"], row["quantity"], row["unit_price"], row["amount"]]
        )
        for cell in ws[ws.max_row]:
            cell.border, cell.alignment = box, wrap
    for label, value in (
        ("Cộng trước thuế", t["amount"]),
        ("Thuế GTGT", t["vat"]),
        ("Tổng cộng", t["total"]),
    ):
        ws.append([None, label, None, None, None, _money(value)])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = box
        ws.cell(row=r, column=2).font = bold
        ws.cell(row=r, column=6).font = bold
    line(f"Bằng chữ: {amount_in_words(t['total'])}.", font=Font(italic=True, bold=True))
    line("Điều 2. Chất lượng dịch vụ", font=bold)
    line(
        "Bên B đã tổ chức các lớp đào tạo theo nội dung hợp đồng. Lịch học, sĩ số, chuyên cần và "
        "kết quả học tập chi tiết theo Bảng kê theo lớp và Danh sách học viên kèm theo biên bản này.",
        height=30,
    )
    line("Điều 3. Thanh toán", font=bold)
    line(
        f"Bên A thanh toán cho Bên B số tiền {_money(t['total'])} đồng "
        f"({amount_in_words(t['total']).lower()}) trong vòng {contract.payment_terms_days} ngày "
        "kể từ ngày nhận được hóa đơn giá trị gia tăng hợp lệ.",
        height=30,
    )
    line("Điều 4. Điều khoản chung", font=bold)
    line("Biên bản được lập thành 04 (bốn) bản có giá trị pháp lý như nhau, mỗi bên giữ 02 (hai) bản.")
    line()
    ws.append(["ĐẠI DIỆN BÊN A", None, None, "ĐẠI DIỆN BÊN B"])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    ws.cell(row=r, column=1).font = ws.cell(row=r, column=4).font = bold
    ws.cell(row=r, column=1).alignment = ws.cell(row=r, column=4).alignment = center
    ws.append(["(Ký, ghi rõ họ tên, đóng dấu)", None, None, "(Ký, ghi rõ họ tên, đóng dấu)"])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    ws.cell(row=r, column=1).alignment = ws.cell(row=r, column=4).alignment = center

    # Sheet 2: per class facts.
    facts = doc["facts"]
    ws2 = book.create_sheet("Bảng kê theo lớp")
    ws2.append([f"BẢNG KÊ CÁC LỚP ĐÀO TẠO – {batch.title} – Hợp đồng {contract.number}"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    ws2.append(
        [
            "STT",
            "Lớp",
            "Khóa học",
            "Từ ngày",
            "Đến ngày",
            "Số buổi",
            "Số giờ",
            "HV danh sách",
            "HV có mặt",
            "HV đạt",
        ]
    )
    for cell in ws2[ws2.max_row]:
        cell.font, cell.alignment, cell.border = bold, center, box
    for i, f in enumerate(facts, start=1):
        c = f["class"]
        ws2.append(
            [
                i,
                c.cohort.name,
                c.cohort.course.title,
                _day(f["first"]),
                _day(f["last"]),
                f["sessions"],
                f["hours"],
                f["listed"],
                f["attended"],
                f["passed"],
            ]
        )
        for cell in ws2[ws2.max_row]:
            cell.border = box
    for col, width in zip("ABCDEFGHIJ", (6, 30, 34, 12, 12, 9, 9, 12, 11, 9), strict=True):
        ws2.column_dimensions[col].width = width

    # Sheet 3: learners (evidence).
    ws3 = book.create_sheet("Danh sách học viên")
    ws3.append(
        [
            "STT",
            "Lớp",
            "Họ tên",
            "Mã NV",
            "Đơn vị",
            "Buổi có mặt / đã học",
            "Chuyên cần (%)",
            "Điểm (%)",
            "Kết quả",
        ]
    )
    for cell in ws3[1]:
        cell.font, cell.alignment, cell.border = bold, center, box
    labels = doc["result_labels"]
    n = 0
    for f in facts:
        for r in f["learners"]:
            n += 1
            ws3.append(
                [
                    n,
                    f["class"].cohort.name,
                    r.get("name", ""),
                    r.get("employee_code", ""),
                    r.get("unit", ""),
                    f"{r.get('sessions_attended', 0)}/{r.get('sessions_taken', 0)}",
                    r.get("attendance_rate"),
                    r.get("grade_percent"),
                    labels.get(r.get("result"), r.get("result", "")),
                ]
            )
            for cell in ws3[ws3.max_row]:
                cell.border = box
    for col, width in zip("ABCDEFGHI", (6, 28, 26, 12, 20, 16, 13, 10, 13), strict=True):
        ws3.column_dimensions[col].width = width
    for sheet in (ws2, ws3):
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0

    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()
