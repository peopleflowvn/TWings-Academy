"""
Excel imports of corporate training: a client's class list (DSHV) into a class, a client's training
plan into a draft plan. Upload -> read -> check every row and preview what will happen -> confirm.

Columns are found by their header, written the way Vietnamese HR files usually write them (accents,
case and punctuation ignored), or by the client's ImportTemplate when its file says it differently.
Nothing is written to the class or plan before confirmation; the original file is kept (private).
"""

import hashlib
import re
import secrets
import unicodedata
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.storage import storages
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone

from apps.catalog.models import Course
from apps.crm.learners import find_learner

from . import services
from .models import ClientUnit, CorporateEmployee, CorporateEnrollment, ImportBatch, PlanLine
from .services import CorporateError

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
HEADER_SCAN_ROWS = 15

# field: (label, header aliases, required)
FIELDS = {
    "roster": {
        "full_name": (
            "Họ tên",
            ["ho ten", "ho va ten", "ten hoc vien", "hoc vien", "full name", "ten nhan vien"],
            True,
        ),
        "email": ("Email", ["email", "e mail", "thu dien tu", "email hoc vien", "email ca nhan"], True),
        "phone": ("Điện thoại", ["dien thoai", "sdt", "so dien thoai", "di dong", "phone", "dt"], False),
        "employee_code": (
            "Mã nhân viên",
            ["ma nv", "ma nhan vien", "ma cbnv", "ma so nhan vien", "ma hr", "msnv"],
            False,
        ),
        "job_title": ("Chức danh", ["chuc danh", "chuc vu", "vi tri", "vi tri cong viec"], False),
        "unit": ("Đơn vị", ["don vi", "phong ban", "chi nhanh", "khoi", "don vi cong tac"], False),
        "manager_name": (
            "CBQL – họ tên",
            ["cbql", "ho ten cbql", "ten cbql", "quan ly truc tiep", "can bo quan ly"],
            False,
        ),
        "manager_email": (
            "CBQL – email",
            ["email cbql", "email quan ly", "email nguoi quan ly", "email quan ly truc tiep"],
            False,
        ),
    },
    "plan": {
        "title": (
            "Nội dung",
            ["noi dung", "noi dung dao tao", "ten khoa hoc", "khoa hoc", "chuong trinh", "ten chuong trinh"],
            True,
        ),
        "unit": ("Đơn vị", ["don vi", "phong ban", "khoi", "doi tuong"], False),
        "planned_start": ("Từ ngày", ["tu ngay", "bat dau", "ngay bat dau", "thoi gian bat dau"], False),
        "planned_end": ("Đến ngày", ["den ngay", "ket thuc", "ngay ket thuc", "thoi gian ket thuc"], False),
        "planned_classes": ("Số lớp", ["so lop"], False),
        "planned_learners": ("Số học viên", ["so hoc vien", "so hv", "so luong", "so luong hoc vien"], False),
        "budget": ("Ngân sách", ["ngan sach", "du toan", "kinh phi", "chi phi du kien"], False),
        "note": ("Ghi chú", ["ghi chu"], False),
    },
}


def normalize(text) -> str:
    """'Mã CBNV (*)' -> 'ma cbnv': no accents, lower case, words only."""
    text = str(text or "").replace("đ", "d").replace("Đ", "D")
    text = "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text.lower()).split())


def match_columns(headers: list, kind: str, mapping: dict | None = None) -> dict[str, int]:
    """{field: column index}: the template's mapping, else exact aliases, else the longest alias contained."""
    found: dict[str, int] = {}
    normalized = [normalize(h) for h in headers]
    if mapping:
        for field, header in mapping.items():
            if field in FIELDS[kind] and normalize(header) in normalized:
                found[field] = normalized.index(normalize(header))
        return found
    taken: set[int] = set()
    for exact in (True, False):
        candidates = []
        for field, (_, aliases, _) in FIELDS[kind].items():
            if field in found:
                continue
            for idx, header in enumerate(normalized):
                if idx in taken or not header:
                    continue
                for alias in aliases:
                    if header == alias if exact else re.search(rf"\b{re.escape(alias)}\b", header):
                        candidates.append((len(alias), field, idx))
        # Longest alias first: "email cbql" is the manager's e-mail, not the learner's.
        for _, field, idx in sorted(candidates, reverse=True):
            if field not in found and idx not in taken:
                found[field] = idx
                taken.add(idx)
    return found


def read_rows(content: bytes, kind: str, template=None) -> tuple[dict, list[tuple[int, dict]]]:
    """({field: header}, [(row number, {field: raw value})]) from the first (or template's) sheet."""
    from io import BytesIO

    from openpyxl import load_workbook

    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable file is the same answer for staff
        raise CorporateError("Không đọc được file. Chỉ nhận file Excel .xlsx.") from exc
    sheet = book[template.sheet] if template and template.sheet in book.sheetnames else book.worksheets[0]
    grid = [list(r) for r in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + HEADER_SCAN_ROWS)]
    book.close()
    if template:
        header_at = template.header_row - 1
    else:  # the header is the row (among the first ones) where most columns are recognised
        scores = [len(match_columns(r, kind)) for r in grid[:HEADER_SCAN_ROWS]]
        header_at = max(range(len(scores)), key=scores.__getitem__) if scores else 0
    if header_at >= len(grid):
        raise CorporateError("File không có dòng tiêu đề.")
    headers = grid[header_at]
    columns = match_columns(headers, kind, template.mapping if template else None)
    missing = [
        FIELDS[kind][f][0] for f, (_, _, required) in FIELDS[kind].items() if required and f not in columns
    ]
    if missing:
        raise CorporateError(f"Không tìm thấy cột: {', '.join(missing)}.")
    rows = []
    for offset, raw in enumerate(grid[header_at + 1 :], start=header_at + 2):
        values = {f: raw[i] if i < len(raw) else None for f, i in columns.items()}
        if all(v in (None, "") for v in values.values()):
            continue
        rows.append((offset, values))
    if len(rows) > MAX_ROWS:
        raise CorporateError(f"Tối đa {MAX_ROWS} dòng mỗi lần nhập.")
    return {f: str(headers[i] or "") for f, i in columns.items()}, rows


# ---------------------------------------------------------------- cell values
def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _phone(value) -> str:
    text = _text(value)
    digits = re.sub(r"\D", "", text)
    if digits and len(digits) == 9 and not text.startswith("0"):  # Excel dropped the leading 0
        return "0" + digits
    return text


def _date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Ngày không hợp lệ: {text}")


def _int(value) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, int | float):
        return int(value)
    digits = re.sub(r"[^\d]", "", str(value))
    if not digits:
        raise ValueError(f"Không phải số: {value}")
    return int(digits)


def _email(value) -> str:
    email = _text(value).lower()
    if email:
        try:
            validate_email(email)
        except ValidationError as exc:
            raise ValueError(f"Email không hợp lệ: {email}") from exc
    return email


# ---------------------------------------------------------------- previews
def _unit_by_name(client) -> dict[str, ClientUnit]:
    return {normalize(u.name): u for u in client.units.all()}


def preview_roster(corporate_class, rows) -> list[dict]:
    units = _unit_by_name(corporate_class.client)
    seen: set[str] = set()
    out = []
    for number, raw in rows:
        errors, values = [], {}
        try:
            values = {
                "full_name": _text(raw.get("full_name")),
                "email": _email(raw.get("email")),
                "phone": _phone(raw.get("phone")),
                "employee_code": _text(raw.get("employee_code")),
                "job_title": _text(raw.get("job_title")),
                "unit": _text(raw.get("unit")),
                "manager_name": _text(raw.get("manager_name")),
                "manager_email": _email(raw.get("manager_email")),
            }
        except ValueError as exc:
            errors.append(str(exc))
        if not errors:
            if not values["full_name"]:
                errors.append("Thiếu họ tên")
            if not values["email"]:
                errors.append("Thiếu email")
            elif values["email"] in seen:
                errors.append("Trùng email với dòng phía trên")
        action, notes = "error", []
        if not errors:
            seen.add(values["email"])
            if values["unit"] and normalize(values["unit"]) not in units:
                notes.append(f"Tạo đơn vị mới “{values['unit']}”")
            action = _roster_action(corporate_class, values, notes)
        out.append({"row": number, "values": values, "action": action, "errors": errors, "notes": notes})
    return out


def _roster_action(corporate_class, values: dict, notes: list) -> str:
    learner = find_learner(email=values["email"], phone=values["phone"])
    if learner is None:
        return "add"
    notes.append(f"Đã có hồ sơ học viên {learner.code}")
    listing = CorporateEnrollment.objects.filter(corporate_class=corporate_class, learner=learner).first()
    if listing is None or not listing.may_study:
        return "add"
    current = CorporateEmployee.objects.filter(
        client=corporate_class.client, learner=learner, valid_to__isnull=True
    ).first()
    compare = ("employee_code", "job_title", "manager_name", "manager_email")
    if current and any(values[f] and values[f] != getattr(current, f) for f in compare):
        return "update"
    return "unchanged"


def preview_plan(plan, rows) -> list[dict]:
    units = _unit_by_name(plan.client)
    courses = {normalize(c.title): c for c in Course.objects.exclude(status="archived")}
    out = []
    for number, raw in rows:
        errors, notes, values = [], [], {}
        try:
            start, end = _date(raw.get("planned_start")), _date(raw.get("planned_end"))
            values = {
                "title": _text(raw.get("title")),
                "unit": _text(raw.get("unit")),
                "planned_start": start.isoformat() if start else None,
                "planned_end": end.isoformat() if end else None,
                "planned_classes": _int(raw.get("planned_classes")) or 1,
                "planned_learners": _int(raw.get("planned_learners")),
                "budget": _int(raw.get("budget")),
                "note": _text(raw.get("note")),
            }
            if start and end and end < start:
                errors.append("Ngày kết thúc trước ngày bắt đầu")
        except ValueError as exc:
            errors.append(str(exc))
        if not errors and not values["title"]:
            errors.append("Thiếu nội dung / tên khóa")
        if not errors:
            course = courses.get(normalize(values["title"]))
            values["course"] = course.pk if course else None
            notes.append(f"Khóa học: {course.title}" if course else "Chưa khớp khóa học nào (chọn sau)")
            if values["unit"] and normalize(values["unit"]) not in units:
                notes.append(f"Tạo đơn vị mới “{values['unit']}”")
        out.append(
            {
                "row": number,
                "values": values,
                "action": "error" if errors else "add",
                "errors": errors,
                "notes": notes,
            }
        )
    return out


def summarize(rows: list[dict]) -> dict:
    counts = {"add": 0, "update": 0, "unchanged": 0, "error": 0}
    for r in rows:
        counts[r["action"]] = counts.get(r["action"], 0) + 1
    return counts


# ---------------------------------------------------------------- batch life cycle
def create_batch(
    *, client, kind: str, upload, user, corporate_class=None, plan=None, template=None
) -> ImportBatch:
    if upload.size > MAX_BYTES:
        raise CorporateError("File tối đa 5 MB.")
    content = upload.read()
    if not content.startswith(b"PK\x03\x04"):
        raise CorporateError("Chỉ nhận file Excel .xlsx.")
    target = {"corporate_class": corporate_class} if kind == "roster" else {"plan": plan}
    if kind == "roster" and corporate_class is None or kind == "plan" and plan is None:
        raise CorporateError("Chọn lớp (danh sách học viên) hoặc kế hoạch để nhập vào.")
    if kind == "plan" and plan.status != "draft":
        raise CorporateError("Chỉ nhập vào kế hoạch đang soạn.")
    if template is not None and (template.client_id != client.pk or template.kind != kind):
        raise CorporateError("Mẫu nhập không thuộc khách hàng / loại file này.")
    fingerprint = hashlib.sha256(content).hexdigest()
    if ImportBatch.objects.filter(fingerprint=fingerprint, status="imported", **target).exists():
        raise CorporateError("File này đã được nhập vào đây rồi.")

    columns, rows = read_rows(content, kind, template)
    preview = preview_roster(corporate_class, rows) if kind == "roster" else preview_plan(plan, rows)
    name = storages["private"].save(
        f"corporate-imports/{client.pk}/{secrets.token_hex(8)}.xlsx", ContentFile(content)
    )
    return ImportBatch.objects.create(
        client=client,
        kind=kind,
        template=template,
        file_name=getattr(upload, "name", "")[:255] or "import.xlsx",
        file_key=name,
        fingerprint=fingerprint,
        columns=columns,
        rows=preview,
        summary=summarize(preview),
        created_by=user,
        **target,
    )


def confirm(batch: ImportBatch) -> ImportBatch:
    """Apply the previewed rows (errors and unchanged rows are skipped). Runs once."""
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
        if batch.status != "previewed":
            raise CorporateError("Lần nhập này đã được xử lý.")
        if batch.kind == "roster":
            done = _apply_roster(batch)
        else:
            done = _apply_plan(batch)
        batch.status, batch.imported_at = "imported", timezone.now()
        batch.summary = {**batch.summary, "applied": done}
        batch.save(update_fields=["status", "imported_at", "summary", "updated_at"])
    return batch


def _unit(client, name: str, units: dict) -> ClientUnit | None:
    if not name:
        return None
    key = normalize(name)
    if key not in units:
        units[key] = ClientUnit.objects.create(
            client=client, kind="department", name=name, parent=services.default_legal_entity(client)
        )
    return units[key]


def _apply_roster(batch: ImportBatch) -> int:
    corporate_class = batch.corporate_class
    client = corporate_class.client
    units = _unit_by_name(client)
    done = 0
    for row in batch.rows:
        if row["action"] not in ("add", "update"):
            continue
        v = row["values"]
        learner = services.find_or_create_learner(
            full_name=v["full_name"], email=v["email"], phone=v["phone"]
        )
        employment = {
            f: v[f] for f in ("employee_code", "job_title", "manager_name", "manager_email") if v[f]
        }
        unit = _unit(client, v["unit"], units)
        if unit is not None:
            employment["unit"] = unit
        services.list_learner(corporate_class, learner, **employment)
        done += 1
    return done


def _apply_plan(batch: ImportBatch) -> int:
    plan = batch.plan
    if plan.status != "draft":
        raise CorporateError("Kế hoạch đã duyệt: không nhập thêm.")
    units = _unit_by_name(plan.client)
    start = plan.lines.count()
    lines = []
    for row in batch.rows:
        if row["action"] != "add":
            continue
        v = row["values"]
        lines.append(
            PlanLine(
                plan=plan,
                course_id=v.get("course"),
                title=v["title"],
                unit=_unit(plan.client, v["unit"], units),
                planned_start=v["planned_start"],
                planned_end=v["planned_end"],
                planned_classes=v["planned_classes"],
                planned_learners=v["planned_learners"],
                budget=v["budget"],
                note=v["note"],
                sort_order=start + len(lines),
            )
        )
    PlanLine.objects.bulk_create(lines)
    return len(lines)
