"""Acceptance (nghiệm thu): quantities from published results × contract rates, approval flow, adjustments."""

from datetime import date, timedelta
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.rbac import Role
from apps.catalog.models import CohortSession
from apps.corporate import services
from apps.corporate.models import AcceptanceLine, ContractRate, ResultSnapshot, ServiceContract

pytestmark = pytest.mark.django_db

URL = "/api/v1/staff/corporate/acceptances/"


@pytest.fixture
def setup(course):
    client = services.create_client(name="Ngân hàng Mẫu", code="NHM")
    contract = ServiceContract.objects.create(client=client, number="HD-01", vat_rate=10, status="active")
    ContractRate.objects.create(contract=contract, kind="per_learner", unit_price=1_000_000)  # general
    ContractRate.objects.create(
        contract=contract, kind="per_learner", course=course, unit_price=2_000_000, vat_rate=8
    )
    ContractRate.objects.create(contract=contract, kind="per_session", unit_price=5_000_000)
    cls = services.create_class(
        client=client,
        course=course,
        name="NHM – Lớp 1",
        cohort_fields={"start_date": date.today()},
        contract=contract,
    )
    now = timezone.now()
    for days in (-3, -1, 5):  # two sessions held, one ahead
        start = now + timedelta(days=days)
        CohortSession.objects.create(cohort=cls.cohort, starts_at=start, ends_at=start + timedelta(hours=3))
    learners = [
        {"enrollment": "a", "name": "An", "sessions_attended": 2, "result": "passed"},
        {"enrollment": "b", "name": "Bình", "sessions_attended": 1, "result": "not_passed"},
        {"enrollment": "c", "name": "Cường", "sessions_attended": 0, "result": "pending"},
    ]
    ResultSnapshot.objects.create(corporate_class=cls, status="published", learners=learners, kpis={})
    return contract, cls


def test_batch_from_results_and_rates(setup, staff_client):
    contract, cls = setup
    corp = staff_client(Role.CORPORATE_TRAINING)
    eligible = corp.get(f"{URL}eligible/", {"contract": contract.pk}).json()
    assert [c["id"] for c in eligible] == [cls.pk]

    res = corp.post(
        URL,
        {"contract": contract.pk, "title": "Đợt 1/2026", "periodTo": date.today().isoformat()},
        format="json",
    )
    assert res.status_code == 201, res.content
    batch = res.json()
    lines = {line["unit"]: line for line in batch["lines"]}
    # The course-specific learner rate beats the general one; VAT from the rate, else the contract.
    assert (lines["học viên"]["quantity"], lines["học viên"]["unitPrice"], lines["học viên"]["amount"]) == (
        3,
        2_000_000,
        6_000_000,
    )
    assert lines["học viên"]["vatAmount"] == 480_000
    assert (lines["buổi"]["quantity"], lines["buổi"]["amount"], lines["buổi"]["vatAmount"]) == (
        2,
        10_000_000,
        1_000_000,
    )
    assert batch["totals"] == {"amount": 16_000_000, "vat": 1_480_000, "total": 17_480_000, "lines": 2}

    assert corp.get(f"{URL}eligible/", {"contract": contract.pk}).json() == []  # billed once
    assert corp.post(URL, {"contract": contract.pk, "title": "Lần 2"}, format="json").status_code == 400


def test_learner_basis_and_hours(setup, staff_client):
    contract, cls = setup
    ContractRate.objects.create(contract=contract, kind="per_hour", unit_price=1_000_000)
    corp = staff_client(Role.CORPORATE_TRAINING)
    batch = corp.post(
        URL, {"contract": contract.pk, "title": "Đạt", "learnerBasis": "passed"}, format="json"
    ).json()
    lines = {line["unit"]: line for line in batch["lines"]}
    assert lines["học viên"]["quantity"] == 1 and lines["giờ"]["quantity"] == 6  # 2 held × 3 h


def test_flow_accept_then_adjust_in_a_new_batch(setup, staff_client):
    contract, cls = setup
    corp = staff_client(Role.CORPORATE_TRAINING)
    batch = corp.post(URL, {"contract": contract.pk, "title": "Đợt 1"}, format="json").json()
    b = f"{URL}{batch['id']}/"
    assert corp.post(f"{b}accept/").status_code == 400  # not approved yet
    assert corp.post(f"{b}submit/").json()["status"] == "submitted"
    assert staff_client(Role.ACADEMIC_MANAGEMENT).post(f"{b}approve/").status_code == 403
    assert corp.post(f"{b}approve/").json()["status"] == "approved"
    accepted = corp.post(
        f"{b}accept/", {"invoiceNumber": "0001234", "invoiceDate": "2026-11-05"}, format="json"
    ).json()
    assert accepted["status"] == "accepted" and accepted["invoiceNumber"] == "0001234"
    learner_line = next(line for line in accepted["lines"] if line["unit"] == "học viên")
    assert (
        corp.post(f"{b}remove-line/", {"line": learner_line["id"]}, format="json").status_code == 400
    )  # final

    fix = corp.post(
        URL, {"contract": contract.pk, "title": "Điều chỉnh đợt 1", "adjustmentOnly": True}, format="json"
    ).json()
    f = f"{URL}{fix['id']}/"
    payload = {
        "adjusts": learner_line["id"],
        "description": "Bớt 1 học viên nghỉ từ buổi đầu",
        "amount": -2_000_000,
    }
    adjusted = corp.post(f"{f}adjustments/", payload, format="json").json()
    line = adjusted["lines"][0]
    assert (line["kind"], line["amount"], line["vatRate"], line["vatAmount"]) == (
        "adjustment",
        -2_000_000,
        8,
        -160_000,
    )
    assert adjusted["totals"]["total"] == -2_160_000
    draft_line = AcceptanceLine.objects.get(pk=line["id"])
    bad = corp.post(
        f"{f}adjustments/", {"adjusts": draft_line.pk, "description": "x", "amount": 1}, format="json"
    )
    assert bad.status_code == 400  # only lines of accepted batches are corrected

    xlsx = corp.get(f"{b}export/")
    values = [
        c for row in load_workbook(BytesIO(xlsx.content)).active.iter_rows(values_only=True) for c in row
    ]
    assert "Tổng cộng" in values and 17_480_000 in values
