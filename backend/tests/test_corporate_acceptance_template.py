"""The default acceptance document (Excel + printable page): amount in words, parties, articles, evidence."""

import pytest

from apps.corporate.acceptance_template import amount_in_words


@pytest.mark.parametrize(
    "amount, words",
    [
        (0, "Không đồng"),
        (5, "Năm đồng"),
        (15, "Mười lăm đồng"),
        (21, "Hai mươi mốt đồng"),
        (24, "Hai mươi tư đồng"),
        (105, "Một trăm linh năm đồng"),
        (110, "Một trăm mười đồng"),
        (1_001, "Một nghìn không trăm linh một đồng"),
        (1_005_000, "Một triệu không trăm linh năm nghìn đồng"),
        (17_480_000, "Mười bảy triệu bốn trăm tám mươi nghìn đồng"),
        (2_000_000_000, "Hai tỷ đồng"),
        (1_000_000_001, "Một tỷ không trăm linh một đồng"),
        (1_200_000_000_000, "Một nghìn hai trăm tỷ đồng"),
        (-2_160_000, "Âm hai triệu một trăm sáu mươi nghìn đồng"),
    ],
)
def test_amount_in_words(amount, words):
    assert amount_in_words(amount) == words


@pytest.mark.django_db
def test_workbook_and_printable_page(course, settings, staff_client):
    from io import BytesIO

    from openpyxl import load_workbook

    from apps.corporate import services
    from apps.corporate.acceptance import build_batch
    from apps.corporate.acceptance_template import export_xlsx
    from apps.corporate.models import ContractRate, ResultSnapshot, ServiceContract

    settings.CORPORATE_PROVIDER = {
        **settings.CORPORATE_PROVIDER,
        "legal_name": "Công ty TNTalent",
        "tax_code": "0101",
    }
    client = services.create_client(name="Ngân hàng Mẫu", code="NHM", address="Hà Nội")
    contract = ServiceContract.objects.create(
        client=client, number="HD-01", vat_rate=8, payment_terms_days=45
    )
    ContractRate.objects.create(contract=contract, kind="per_learner", unit_price=1_000_000)
    cls = services.create_class(
        client=client, course=course, name="NHM – Lớp 1", cohort_fields={}, contract=contract
    )
    learners = [
        {
            "enrollment": "a",
            "name": "Trần Thị Bình",
            "employee_code": "NV01",
            "unit": "Bán lẻ",
            "sessions_attended": 3,
            "sessions_taken": 4,
            "attendance_rate": 75,
            "result": "passed",
        }
    ]
    ResultSnapshot.objects.create(corporate_class=cls, status="published", learners=learners)
    batch = build_batch(
        contract, title="Đợt 1", period_from=None, period_to=None, learner_basis="listed", user=None
    )

    book = load_workbook(BytesIO(export_xlsx(batch)))
    assert book.sheetnames == ["Biên bản nghiệm thu", "Bảng kê theo lớp", "Danh sách học viên"]
    text = " | ".join(
        str(c) for row in book["Biên bản nghiệm thu"].iter_rows(values_only=True) for c in row if c
    )
    assert "CÔNG TY TNTALENT" in text and "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in text
    assert "BÊN A (Bên sử dụng dịch vụ): Ngân hàng Mẫu" in text and "Địa chỉ: Hà Nội" in text
    assert "BÊN B (Bên cung cấp dịch vụ): Công ty TNTalent" in text and "Mã số thuế: 0101" in text
    assert "1.080.000" in text and "Bằng chữ: Một triệu không trăm tám mươi nghìn đồng." in text
    assert "trong vòng 45 ngày" in text and "Đại diện: ……" in text  # blanks to fill in by hand
    rows = list(book["Danh sách học viên"].iter_rows(values_only=True))
    assert rows[1][2] == "Trần Thị Bình" and rows[1][8] == "Đạt"

    # The same minutes as a printable page (In / Lưu PDF), for staff who may see the client only.
    from apps.accounts.rbac import Role

    url = f"/api/v1/staff/corporate/acceptances/{batch.pk}/print/"
    page = staff_client(Role.CORPORATE_TRAINING).get(url)
    html = page.content.decode()
    assert page.status_code == 200 and page["Content-Type"].startswith("text/html")
    assert "BIÊN BẢN NGHIỆM THU DỊCH VỤ ĐÀO TẠO" in html and "CÔNG TY TNTALENT" in html
    assert "Bằng chữ: Một triệu không trăm tám mươi nghìn đồng." in html and "trong vòng 45 ngày" in html
    assert "Trần Thị Bình" in html and "NV01" in html and "Đạt" in html  # learner list
    assert staff_client(Role.SALES_CRM).get(url).status_code == 403
